from datetime import datetime
from io import BytesIO
import re
from uuid import uuid4
from urllib.parse import urlparse

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from fastapi.responses import Response
from sqlalchemy.orm import Session
from docx import Document as DocxDocument

from .. import database, models, schemas
from ..services.volcengine_knowledge_base import (
	add_document_by_uri,
	create_remote_collection,
	delete_document,
	delete_remote_collection,
	format_search_hits,
	format_timestamp,
	get_document_info,
	infer_doc_type,
	list_document_points,
	list_documents,
	search_knowledge,
)
from ..services.tos_storage import build_knowledge_base_prefix, delete_tos_object_by_uri, delete_tos_prefix, parse_tos_uri, upload_bytes_to_tos
from ..services.tos_storage import get_tos_object_bytes_by_uri


router = APIRouter(prefix="/knowledge-bases", tags=["knowledge-bases"])


def _serialize_document(document: models.KnowledgeDocument) -> schemas.KnowledgeDocumentSummary:
	remote_meta = document.remote_meta or {}
	status_payload = remote_meta.get("status") if isinstance(remote_meta, dict) else None
	failed_code = None
	remote_process_status = None
	if isinstance(status_payload, dict):
		failed_code = status_payload.get("failed_code")
		remote_process_status = status_payload.get("process_status")
	serialized_chunks = []
	for chunk in document.chunks or []:
		if isinstance(chunk, dict):
			content = str(chunk.get("content") or "").strip()
			attachment_link = str(chunk.get("attachment_link") or "").strip() or None
			if not content and not attachment_link:
				continue
			serialized_chunks.append(
				schemas.KnowledgeDocumentChunk(
					id=str(chunk.get("id") or chunk.get("point_id") or "").strip() or None,
					point_id=str(chunk.get("point_id") or "").strip() or None,
					chunk_id=int(chunk.get("chunk_id")) if chunk.get("chunk_id") is not None else None,
					content=content,
					chunk_type=str(chunk.get("chunk_type") or "").strip() or None,
					attachment_link=attachment_link,
					page_numbers=[int(page) for page in (chunk.get("page_numbers") or []) if str(page).strip()],
					original_coordinate=chunk.get("original_coordinate"),
					doc_id=str(chunk.get("doc_id") or "").strip() or None,
					doc_name=str(chunk.get("doc_name") or "").strip() or None,
				)
			)
			continue
		chunk_text = str(chunk or "").strip()
		if chunk_text:
			serialized_chunks.append(chunk_text)
	return schemas.KnowledgeDocumentSummary(
		id=document.id,
		name=document.name,
		fileType=(document.file_type or "TXT").upper(),
		updatedAt=document.updated_at_display or format_timestamp(document.updated_at or document.created_at),
		status=document.status or "pending",
		summary=document.summary or "",
		chunks=serialized_chunks,
		remote_doc_id=document.remote_doc_id,
		source_url=document.source_url,
		size_bytes=int(document.size_bytes or 0),
		remote_process_status=remote_process_status if isinstance(remote_process_status, int) else None,
		failed_code=failed_code if isinstance(failed_code, int) else None,
	)


def _serialize_knowledge_base(item: models.KnowledgeBase) -> schemas.KnowledgeBaseCreateResponse:
	sync_info = item.sync_info or None
	sync_info_schema = None
	if isinstance(sync_info, dict) and sync_info.get("resource_id"):
		sync_info_schema = schemas.KnowledgeBaseSyncInfo(**sync_info)
	return schemas.KnowledgeBaseCreateResponse(
		id=item.id,
		name=item.name,
		description=item.description or "",
		date=item.date or format_timestamp(item.created_at),
		enabled=bool(item.enabled),
		status=bool(item.status),
		type=item.type or "text",
		teams=[str(team) for team in (item.teams or []) if str(team).strip()],
		documents=[_serialize_document(document) for document in item.documents or []],
		remote_provider=item.remote_provider or "volcengine",
		remote_resource_id=item.remote_resource_id,
		local_display_name=item.local_display_name,
		remote_host=item.remote_host,
		remote_project=item.remote_project,
		remote_collection_name=item.remote_collection_name,
		sync_status=item.sync_status or "pending",
		sync_error=item.sync_error,
		sync_info=sync_info_schema,
	)


def _default_doc_name_from_uri(uri: str) -> str:
	parsed = urlparse(str(uri or "").strip())
	path = (parsed.path or "").strip()
	if path and "/" in path:
		candidate = path.rsplit("/", 1)[-1]
		if candidate:
			return candidate
	return path or "document"


def _extract_markdown_from_docx_bytes(content: bytes) -> str:
	document = DocxDocument(BytesIO(content))
	lines: list[str] = []
	for paragraph in document.paragraphs:
		text = str(paragraph.text or "").strip()
		if not text:
			continue
		style_name = str(getattr(getattr(paragraph, "style", None), "name", "") or "").lower()
		if style_name.startswith("heading"):
			level_digits = "".join(ch for ch in style_name if ch.isdigit())
			level = max(1, min(int(level_digits or "1"), 6))
			lines.append(f"{'#' * level} {text}")
		else:
			lines.append(text)

	for table in document.tables:
		for row in table.rows:
			cells = [str(cell.text or "").strip() for cell in row.cells]
			if any(cells):
				lines.append(" | ".join(cells))

	return "\n\n".join(lines).strip()


def _extract_text_preview(file_type: str, content: bytes) -> str:
	resolved_file_type = str(file_type or "").strip().lower()
	if resolved_file_type in {"docx"}:
		return _extract_markdown_from_docx_bytes(content)
	if resolved_file_type in {"txt", "md", "markdown", "csv", "json"}:
		return content.decode("utf-8", errors="ignore").strip()
	if resolved_file_type in {"doc"}:
		raise RuntimeError("Legacy .doc preview is not supported yet. Please convert the file to .docx for inline preview.")
	if resolved_file_type in {"ppt", "pptx", "xls", "xlsx"}:
		raise RuntimeError("This file type does not support inline text preview yet.")
	return content.decode("utf-8", errors="ignore").strip()


def _build_safe_inline_filename(file_name: str | None) -> str:
	name = str(file_name or "document").strip() or "document"
	safe_name = re.sub(r"[^A-Za-z0-9._-]", "_", name)
	safe_name = re.sub(r"_+", "_", safe_name).strip("._") or "document"
	return safe_name


def _get_knowledge_base_or_404(knowledge_base_id: str, db: Session) -> models.KnowledgeBase:
	item = db.query(models.KnowledgeBase).filter(models.KnowledgeBase.id == knowledge_base_id).first()
	if item is None:
		raise HTTPException(status_code=404, detail="Knowledge base not found")
	return item


def _persist_imported_document(
	*,
	item: models.KnowledgeBase,
	db: Session,
	remote_result: dict,
	document_name: str,
	document_type: str | None,
	source_url: str,
	tags: list[dict],
	size_bytes: int = 0,
) -> models.KnowledgeDocument:
	document = models.KnowledgeDocument(
		id=f"doc_{uuid4().hex[:16]}",
		knowledge_base_id=item.id,
		remote_doc_id=remote_result.get("doc_id"),
		name=document_name,
		file_type=(document_type or infer_doc_type(doc_name=document_name, uri=source_url) or "txt").upper(),
		updated_at_display=format_timestamp(None),
		status="waiting",
		summary="Document import accepted by Volcengine and is waiting to enter indexing.",
		chunks=[],
		source_url=source_url,
		size_bytes=int(size_bytes or 0),
		remote_meta={
			"request_id": remote_result.get("request_id"),
			"resource_id": remote_result.get("resource_id"),
			"collection_name": remote_result.get("collection_name"),
			"project": remote_result.get("project"),
			"uri": remote_result.get("uri"),
			"tag_list": tags,
			"status": {"process_status": None},
		},
	)
	db.add(document)
	item.date = format_timestamp(None)
	item.sync_status = "indexing"
	item.sync_error = None
	sync_info = dict(item.sync_info or {})
	if sync_info:
		sync_info["status"] = item.sync_status
		sync_info["request_id"] = remote_result.get("request_id")
		item.sync_info = sync_info
	db.commit()
	db.refresh(document)
	return document


def _timestamp_from_remote(value) -> str:
	try:
		if value is None:
			return format_timestamp(None)
		int_value = int(value)
		if int_value > 10_000_000_000:
			int_value = int_value / 1000
		return format_timestamp(datetime.fromtimestamp(int_value))
	except Exception:
		return format_timestamp(None)


def _map_remote_process_status(remote_document: dict | None, local_document: models.KnowledgeDocument) -> tuple[str, str, dict]:
	if not remote_document:
		return "waiting", "Document is waiting to appear in Volcengine document status list.", {"process_status": None}

	status_payload = remote_document.get("status") or {}
	process_status = status_payload.get("process_status")
	failed_code = status_payload.get("failed_code")
	if process_status == 0:
		return "processed", "Document has been indexed successfully and is ready for retrieval.", status_payload
	if process_status == 1:
		if failed_code is not None:
			return "failed", f"Volcengine document processing failed with code {failed_code}.", status_payload
		return "processing", "Document is currently being parsed and indexed by Volcengine.", status_payload
	return "processing", "Document is currently being parsed and indexed by Volcengine.", status_payload


def _extract_page_numbers_from_point(point: dict | None) -> list[int]:
	if not isinstance(point, dict):
		return []
	page_numbers = []
	original_coordinate = point.get("original_coordinate") or {}
	for raw_page_no in (original_coordinate.get("page_no") or []):
		try:
			page_numbers.append(int(raw_page_no) + 1)
		except Exception:
			continue
	return sorted(set(page_numbers))


def _extract_chunk_text_from_point(point: dict | None) -> str:
	if not isinstance(point, dict):
		return ""
	content = str(point.get("content") or "").strip()
	content = re.sub(r"</?KBImage>", "", content, flags=re.IGNORECASE)
	content = re.sub(r"</?KBTable>", "", content, flags=re.IGNORECASE)
	content = content.replace("\u200b", "").replace("\ufeff", "")
	content = re.sub(r"[ \t]+", " ", content)
	content = re.sub(r"\n{3,}", "\n\n", content).strip()
	if content:
		return content
	structured_fields = point.get("table_chunk_fields") or []
	parts = []
	for field in structured_fields:
		if not isinstance(field, dict):
			continue
		field_name = str(field.get("field_name") or "").strip()
		field_value = str(field.get("field_value") or "").strip()
		if field_name and field_value:
			parts.append(f"{field_name}: {field_value}")
		elif field_value:
			parts.append(field_value)
	return "\n".join(parts).strip()


def _serialize_chunk_from_point(point: dict | None, fallback_index: int = 0) -> dict | None:
	if not isinstance(point, dict):
		return None
	content = _extract_chunk_text_from_point(point)
	attachment_link = str(
		point.get("attachment_link")
		or point.get("image_attachment_link")
		or point.get("image_url")
		or point.get("attachment_url")
		or ""
	).strip() or None
	if not content and not attachment_link:
		return None
	doc_info = point.get("doc_info") or {}
	point_id = str(point.get("point_id") or "").strip() or f"point_{fallback_index}"
	return {
		"id": point_id,
		"point_id": point_id,
		"chunk_id": point.get("chunk_id"),
		"content": content,
		"chunk_type": str(point.get("chunk_type") or "").strip().lower() or None,
		"attachment_link": attachment_link,
		"page_numbers": _extract_page_numbers_from_point(point),
		"original_coordinate": point.get("original_coordinate"),
		"doc_id": str(doc_info.get("doc_id") or "").strip() or None,
		"doc_name": str(doc_info.get("doc_name") or doc_info.get("title") or "").strip() or None,
	}


def _fetch_remote_document_chunks(item: models.KnowledgeBase, remote_doc_id: str | None) -> tuple[list[dict], dict | None]:
	resolved_remote_doc_id = str(remote_doc_id or "").strip()
	if not resolved_remote_doc_id:
		return [], None
	point_payload = list_document_points(
		resource_id=item.remote_resource_id,
		collection_name=item.remote_collection_name,
		project=item.remote_project,
		host=item.remote_host,
		doc_ids=[resolved_remote_doc_id],
		get_attachment_link=True,
	)
	chunks: list[dict] = []
	seen = set()
	for index, point in enumerate(point_payload.get("point_list") or []):
		chunk_payload = _serialize_chunk_from_point(point, fallback_index=index)
		if not chunk_payload:
			continue
		chunk_key = chunk_payload.get("point_id") or f"{chunk_payload.get('chunk_id')}::{chunk_payload.get('content')}"
		if chunk_key in seen:
			continue
		seen.add(chunk_key)
		chunks.append(chunk_payload)
	return chunks, point_payload


def _sync_knowledge_base_documents(item: models.KnowledgeBase, db: Session) -> models.KnowledgeBase:
	if not item.remote_resource_id and not item.remote_collection_name:
		return item

	remote_documents_by_id: dict[str, dict] = {}
	try:
		remote_payload = list_documents(
			resource_id=item.remote_resource_id,
			collection_name=item.remote_collection_name,
			project=item.remote_project,
			host=item.remote_host,
		)
		for remote_document in remote_payload.get("doc_list") or []:
			remote_doc_id = str(remote_document.get("doc_id") or "").strip()
			if remote_doc_id:
				remote_documents_by_id[remote_doc_id] = remote_document
	except Exception as exc:
		item.sync_status = "document_sync_error"
		item.sync_error = str(exc)
		sync_info = dict(item.sync_info or {})
		sync_info["status"] = item.sync_status
		item.sync_info = sync_info or None
		db.commit()
		return item

	has_waiting = False
	has_processing = False
	has_failed = False
	has_processed = False

	for document in item.documents or []:
		remote_document = remote_documents_by_id.get(str(document.remote_doc_id or "").strip())
		status, summary, status_payload = _map_remote_process_status(remote_document, document)
		document.status = status
		document.summary = summary
		if remote_document:
			document.name = str(remote_document.get("doc_name") or document.name)
			document.file_type = str(remote_document.get("doc_type") or document.file_type or "TXT").upper()
			document.updated_at_display = _timestamp_from_remote(remote_document.get("update_time") or remote_document.get("create_time"))
			document.size_bytes = int(remote_document.get("doc_size") or document.size_bytes or 0)
			if remote_document.get("url"):
				document.source_url = str(remote_document.get("url"))
			elif remote_document.get("tos_path"):
				document.source_url = f"tos://{remote_document.get('tos_path')}"
			remote_meta = dict(document.remote_meta or {})
			remote_meta.update(
				{
					"status": status_payload,
					"doc_summary": remote_document.get("doc_summary"),
					"brief_summary": remote_document.get("brief_summary"),
					"statistics": remote_document.get("statistics"),
					"raw": remote_document,
				}
			)
			document.remote_meta = remote_meta
			if status == "processed":
				try:
					chunks, point_payload = _fetch_remote_document_chunks(item, document.remote_doc_id)
					if chunks:
						document.chunks = chunks
						remote_meta = dict(document.remote_meta or {})
						remote_meta["point_count"] = len(point_payload.get("point_list") or []) if point_payload else len(chunks)
						remote_meta["point_request_id"] = point_payload.get("request_id") if point_payload else None
						document.remote_meta = remote_meta
				except Exception as exc:
					remote_meta = dict(document.remote_meta or {})
					remote_meta["chunk_sync_error"] = str(exc)
					document.remote_meta = remote_meta

		has_waiting = has_waiting or status == "waiting"
		has_processing = has_processing or status == "processing"
		has_failed = has_failed or status == "failed"
		has_processed = has_processed or status == "processed"

	if has_waiting or has_processing:
		item.sync_status = "indexing"
	elif has_failed and not has_processed:
		item.sync_status = "failed"
	elif has_failed and has_processed:
		item.sync_status = "partial_failed"
	elif has_processed:
		item.sync_status = "connected"
	else:
		item.sync_status = item.sync_status or "pending"
	item.sync_error = None
	item.date = format_timestamp(None)
	sync_info = dict(item.sync_info or {})
	if sync_info:
		sync_info["status"] = item.sync_status
		item.sync_info = sync_info
	db.commit()
	db.refresh(item)
	return item


def _refresh_sync_status_from_local_documents(item: models.KnowledgeBase) -> None:
	documents = item.documents or []
	statuses = {str(document.status or "").lower() for document in documents}
	if "waiting" in statuses or "processing" in statuses or "pending" in statuses:
		item.sync_status = "indexing"
	elif "failed" in statuses and "processed" in statuses:
		item.sync_status = "partial_failed"
	elif "failed" in statuses:
		item.sync_status = "failed"
	elif "processed" in statuses or not documents:
		item.sync_status = "connected"
	else:
		item.sync_status = item.sync_status or "pending"


def _delete_tos_file_for_document(document: models.KnowledgeDocument) -> None:
	uri_candidates = []
	if document.source_url:
		uri_candidates.append(str(document.source_url).strip())
	remote_meta = document.remote_meta or {}
	if isinstance(remote_meta, dict) and remote_meta.get("uri"):
		uri_candidates.append(str(remote_meta.get("uri")).strip())

	for uri in uri_candidates:
		if not uri.startswith("tos://"):
			continue
		try:
			delete_tos_object_by_uri(uri)
		except Exception as exc:
			message = str(exc).lower()
			if "not found" in message or "no such key" in message or "404" in message:
				continue
			raise


def _delete_tos_files_for_knowledge_base(item: models.KnowledgeBase) -> int:
	uri_candidates: set[str] = set()
	prefix_targets: set[tuple[str | None, str]] = {(None, build_knowledge_base_prefix(item.id))}

	for document in item.documents or []:
		if document.source_url:
			uri_candidates.add(str(document.source_url).strip())
		remote_meta = document.remote_meta or {}
		if isinstance(remote_meta, dict) and remote_meta.get("uri"):
			uri_candidates.add(str(remote_meta.get("uri")).strip())

	for uri in uri_candidates:
		if not uri.startswith("tos://"):
			continue

		try:
			delete_tos_object_by_uri(uri)
		except Exception as exc:
			message = str(exc).lower()
			if "not found" not in message and "no such key" not in message and "404" not in message:
				raise

		try:
			bucket_name, object_key = parse_tos_uri(uri)
		except Exception:
			continue

		marker = f"/{item.id}/"
		marker_index = object_key.find(marker)
		if marker_index >= 0:
			inferred_prefix = object_key[: marker_index + len(marker)]
			if inferred_prefix.strip():
				prefix_targets.add((bucket_name, inferred_prefix))

	deleted_count = 0
	for bucket_name, prefix in prefix_targets:
		try:
			deleted_count += delete_tos_prefix(prefix=prefix, bucket_name=bucket_name)
		except Exception as exc:
			message = str(exc).lower()
			if "not found" not in message and "no such key" not in message and "404" not in message:
				raise

	return deleted_count


@router.get("/", response_model=list[schemas.KnowledgeBaseCreateResponse])
def list_knowledge_bases(db: Session = Depends(database.get_db)):
	items = db.query(models.KnowledgeBase).order_by(models.KnowledgeBase.updated_at.desc()).all()
	return [_serialize_knowledge_base(item) for item in items]


@router.get("/{knowledge_base_id}", response_model=schemas.KnowledgeBaseCreateResponse)
def get_knowledge_base(knowledge_base_id: str, db: Session = Depends(database.get_db)):
	item = db.query(models.KnowledgeBase).filter(models.KnowledgeBase.id == knowledge_base_id).first()
	if item is None:
		raise HTTPException(status_code=404, detail="Knowledge base not found")
	return _serialize_knowledge_base(item)


@router.get("/{knowledge_base_id}/documents", response_model=schemas.KnowledgeBaseDocumentsResponse)
def list_knowledge_base_documents(knowledge_base_id: str, refresh: bool = False, db: Session = Depends(database.get_db)):
	item = _get_knowledge_base_or_404(knowledge_base_id, db)
	if refresh:
		item = _sync_knowledge_base_documents(item, db)
	return schemas.KnowledgeBaseDocumentsResponse(
		knowledge_base_id=item.id,
		sync_status=item.sync_status or "pending",
		sync_error=item.sync_error,
		documents=[_serialize_document(document) for document in item.documents or []],
	)


@router.post("/{knowledge_base_id}/documents/sync", response_model=schemas.KnowledgeBaseCreateResponse)
def sync_knowledge_base_documents(knowledge_base_id: str, db: Session = Depends(database.get_db)):
	item = _get_knowledge_base_or_404(knowledge_base_id, db)
	item = _sync_knowledge_base_documents(item, db)
	return _serialize_knowledge_base(item)


@router.get("/{knowledge_base_id}/documents/{document_id}", response_model=schemas.KnowledgeDocumentSummary)
def get_knowledge_base_document(knowledge_base_id: str, document_id: str, refresh: bool = False, db: Session = Depends(database.get_db)):
	item = _get_knowledge_base_or_404(knowledge_base_id, db)
	document = db.query(models.KnowledgeDocument).filter(
		models.KnowledgeDocument.id == document_id,
		models.KnowledgeDocument.knowledge_base_id == knowledge_base_id,
	).first()
	if document is None:
		raise HTTPException(status_code=404, detail="Knowledge base document not found")
	if refresh and document.remote_doc_id and (item.remote_resource_id or item.remote_collection_name):
		try:
			remote_info = get_document_info(
				resource_id=item.remote_resource_id,
				collection_name=item.remote_collection_name,
				project=item.remote_project,
				host=item.remote_host,
				doc_id=document.remote_doc_id,
			)
			remote_document = remote_info.get("data") or {}
			status, summary, status_payload = _map_remote_process_status(remote_document, document)
			document.status = status
			document.summary = summary
			document.name = str(remote_document.get("doc_name") or document.name)
			document.file_type = str(remote_document.get("doc_type") or document.file_type or "TXT").upper()
			document.updated_at_display = _timestamp_from_remote(remote_document.get("update_time") or remote_document.get("create_time"))
			remote_meta = dict(document.remote_meta or {})
			remote_meta.update(
				{
					"status": status_payload,
					"doc_summary": remote_document.get("doc_summary"),
					"brief_summary": remote_document.get("brief_summary"),
					"statistics": remote_document.get("statistics"),
					"raw": remote_document,
				}
			)
			if status == "processed":
				try:
					chunks, point_payload = _fetch_remote_document_chunks(item, document.remote_doc_id)
					if chunks:
						document.chunks = chunks
						remote_meta["point_count"] = len(point_payload.get("point_list") or []) if point_payload else len(chunks)
						remote_meta["point_request_id"] = point_payload.get("request_id") if point_payload else None
				except Exception as exc:
					remote_meta["chunk_sync_error"] = str(exc)
			document.remote_meta = remote_meta
			db.commit()
			db.refresh(document)
		except Exception as exc:
			raise HTTPException(status_code=502, detail=str(exc)) from exc
	return _serialize_document(document)


@router.get("/{knowledge_base_id}/documents/{document_id}/preview/raw")
def preview_knowledge_base_document_raw(knowledge_base_id: str, document_id: str, db: Session = Depends(database.get_db)):
	_get_knowledge_base_or_404(knowledge_base_id, db)
	document = db.query(models.KnowledgeDocument).filter(
		models.KnowledgeDocument.id == document_id,
		models.KnowledgeDocument.knowledge_base_id == knowledge_base_id,
	).first()
	if document is None:
		raise HTTPException(status_code=404, detail="Knowledge base document not found")
	if not document.source_url or not str(document.source_url).startswith("tos://"):
		raise HTTPException(status_code=400, detail="This document does not have a previewable TOS source file")

	try:
		content, content_type, _bucket_name, object_key = get_tos_object_bytes_by_uri(document.source_url)
	except Exception as exc:
		raise HTTPException(status_code=502, detail=f"Failed to fetch preview file from TOS: {exc}") from exc

	file_name = _build_safe_inline_filename(document.name or _default_doc_name_from_uri(object_key))
	headers = {"Content-Disposition": f'inline; filename="{file_name}"'}
	return Response(content=content, media_type=content_type or "application/octet-stream", headers=headers)


@router.get("/{knowledge_base_id}/documents/{document_id}/preview/text", response_model=schemas.KnowledgeDocumentPreviewTextResponse)
def preview_knowledge_base_document_text(knowledge_base_id: str, document_id: str, db: Session = Depends(database.get_db)):
	_get_knowledge_base_or_404(knowledge_base_id, db)
	document = db.query(models.KnowledgeDocument).filter(
		models.KnowledgeDocument.id == document_id,
		models.KnowledgeDocument.knowledge_base_id == knowledge_base_id,
	).first()
	if document is None:
		raise HTTPException(status_code=404, detail="Knowledge base document not found")
	if not document.source_url or not str(document.source_url).startswith("tos://"):
		raise HTTPException(status_code=400, detail="This document does not have a previewable TOS source file")

	try:
		content, _content_type, _bucket_name, _object_key = get_tos_object_bytes_by_uri(document.source_url)
		markdown = _extract_text_preview(document.file_type, content)
	except Exception as exc:
		raise HTTPException(status_code=502, detail=f"Failed to build text preview: {exc}") from exc

	return schemas.KnowledgeDocumentPreviewTextResponse(
		knowledge_base_id=knowledge_base_id,
		document_id=document_id,
		file_type=(document.file_type or "TXT").upper(),
		preview_type="markdown",
		markdown=markdown,
	)


@router.delete("/{knowledge_base_id}/documents/{document_id}", response_model=schemas.KnowledgeBaseCreateResponse)
def delete_knowledge_base_document(knowledge_base_id: str, document_id: str, db: Session = Depends(database.get_db)):
	item = _get_knowledge_base_or_404(knowledge_base_id, db)
	document = db.query(models.KnowledgeDocument).filter(
		models.KnowledgeDocument.id == document_id,
		models.KnowledgeDocument.knowledge_base_id == knowledge_base_id,
	).first()
	if document is None:
		raise HTTPException(status_code=404, detail="Knowledge base document not found")

	if document.remote_doc_id and (item.remote_resource_id or item.remote_collection_name):
		try:
			delete_document(
				resource_id=item.remote_resource_id,
				collection_name=item.remote_collection_name,
				project=item.remote_project,
				host=item.remote_host,
				doc_id=document.remote_doc_id,
			)
		except Exception as exc:
			message = str(exc)
			if "doc not exist" not in message.lower():
				raise HTTPException(status_code=502, detail=message) from exc

	try:
		_delete_tos_file_for_document(document)
	except Exception as exc:
		raise HTTPException(status_code=502, detail=f"Failed to delete TOS object: {exc}") from exc

	db.delete(document)
	item.date = format_timestamp(None)
	item.sync_error = None
	_refresh_sync_status_from_local_documents(item)
	sync_info = dict(item.sync_info or {})
	if sync_info:
		sync_info["status"] = item.sync_status
		item.sync_info = sync_info
	db.commit()
	db.refresh(item)
	return _serialize_knowledge_base(item)


@router.post("/", response_model=schemas.KnowledgeBaseCreateResponse)
def create_knowledge_base(payload: schemas.KnowledgeBaseCreateRequest, db: Session = Depends(database.get_db)):
	item_id = f"kb_{uuid4().hex[:12]}"
	remote_resource_id = (payload.remote_resource_id or "").strip() or None
	remote_host = (payload.remote_host or "").strip() or None
	remote_project = (payload.remote_project or "").strip() or None
	requested_remote_collection_name = (payload.remote_collection_name or "").strip() or None
	remote_collection_name = requested_remote_collection_name
	sync_status = "connected" if remote_resource_id else "pending"
	sync_error = None
	sync_info = None

	if payload.auto_create_remote and not remote_resource_id:
		try:
			remote_info = create_remote_collection(
				name=remote_collection_name or payload.name,
				description=payload.description or "",
				project=remote_project,
				host=remote_host,
			)
			remote_resource_id = remote_info.get("resource_id") or None
			remote_host = remote_info.get("host") or remote_host
			remote_project = remote_info.get("project") or remote_project
			remote_collection_name = remote_info.get("collection_name") or remote_collection_name
			sync_status = remote_info.get("status") or "synced"
			sync_info = {
				"provider": "volcengine",
				"status": sync_status,
				"resource_id": remote_resource_id or "",
				"collection_name": remote_collection_name or payload.name,
				"host": remote_host or "",
				"project": remote_project or "default",
				"request_id": remote_info.get("request_id"),
			}
		except Exception as exc:
			message = str(exc)
			lower_message = message.lower()
			if "collection[" in lower_message and "exist" in lower_message:
				raise HTTPException(
					status_code=409,
					detail=(
						f"Failed to create remote Volcengine knowledge base because collection '{remote_collection_name or payload.name}' already exists. "
						"Use a different knowledge base name or switch to connect an existing remote knowledge base."
					),
				) from exc
			raise HTTPException(
				status_code=502,
				detail=f"Failed to create remote Volcengine knowledge base: {message}",
			) from exc

	if remote_resource_id and not remote_collection_name:
		remote_collection_name = payload.name.strip()

	if not remote_resource_id and not payload.auto_create_remote:
		remote_host = None
		remote_project = None
		remote_collection_name = None

	if remote_resource_id and sync_info is None:
		sync_info = {
			"provider": "volcengine",
			"status": sync_status,
			"resource_id": remote_resource_id,
			"collection_name": remote_collection_name or payload.name,
			"host": remote_host or "",
			"project": remote_project or "default",
			"request_id": None,
		}

	item = models.KnowledgeBase(
		id=item_id,
		name=payload.name,
		description=payload.description or "",
		date=format_timestamp(None),
		enabled=1,
		status=1,
		type=payload.type or "text",
		teams=[str(team) for team in payload.teams if str(team).strip()],
		remote_provider="volcengine",
		remote_resource_id=remote_resource_id,
		local_display_name=(payload.local_display_name or "").strip() or None,
		remote_host=remote_host,
		remote_project=remote_project,
		remote_collection_name=remote_collection_name,
		sync_status=sync_status,
		sync_error=sync_error,
		sync_info=sync_info,
	)
	db.add(item)
	db.commit()
	db.refresh(item)
	return _serialize_knowledge_base(item)


@router.put("/{knowledge_base_id}", response_model=schemas.KnowledgeBaseCreateResponse)
def update_knowledge_base(knowledge_base_id: str, payload: schemas.KnowledgeBaseUpdateRequest, db: Session = Depends(database.get_db)):
	item = db.query(models.KnowledgeBase).filter(models.KnowledgeBase.id == knowledge_base_id).first()
	if item is None:
		raise HTTPException(status_code=404, detail="Knowledge base not found")

	updates = payload.model_dump(exclude_unset=True)
	for key, value in updates.items():
		if key in {"enabled", "status"} and value is not None:
			setattr(item, key, 1 if bool(value) else 0)
		elif key == "teams" and value is not None:
			setattr(item, key, [str(team) for team in value if str(team).strip()])
		else:
			setattr(item, key, value)

	if not item.remote_resource_id:
		item.remote_host = None
		item.remote_project = None
		item.remote_collection_name = None
		if item.sync_status == "connected":
			item.sync_status = "pending"

	sync_info = dict(item.sync_info or {})
	if item.remote_resource_id:
		sync_info.update(
			{
				"provider": "volcengine",
				"status": item.sync_status or "connected",
				"resource_id": item.remote_resource_id,
				"collection_name": item.remote_collection_name or item.name,
				"host": item.remote_host or "",
				"project": item.remote_project or "default",
				"request_id": sync_info.get("request_id"),
			}
		)
	item.sync_info = sync_info or None
	db.commit()
	db.refresh(item)
	return _serialize_knowledge_base(item)


@router.delete("/{knowledge_base_id}")
def delete_knowledge_base(knowledge_base_id: str, db: Session = Depends(database.get_db)):
	item = db.query(models.KnowledgeBase).filter(models.KnowledgeBase.id == knowledge_base_id).first()
	if item is None:
		raise HTTPException(status_code=404, detail="Knowledge base not found")

	if item.remote_resource_id or item.remote_collection_name:
		try:
			delete_remote_collection(
				resource_id=item.remote_resource_id,
				collection_name=item.remote_collection_name,
				project=item.remote_project,
				host=item.remote_host,
			)
		except Exception as exc:
			message = str(exc)
			if "collection not exist" not in message.lower():
				raise HTTPException(status_code=502, detail=f"Failed to delete remote Volcengine knowledge base: {message}") from exc

	try:
		_delete_tos_files_for_knowledge_base(item)
	except Exception as exc:
		raise HTTPException(status_code=502, detail=f"Failed to delete TOS files for knowledge base: {exc}") from exc

	db.delete(item)
	db.commit()
	return {"ok": True, "knowledge_base_id": knowledge_base_id}


@router.post("/{knowledge_base_id}/search", response_model=schemas.KnowledgeBaseSearchResponse)
def debug_search_knowledge_base(knowledge_base_id: str, payload: schemas.KnowledgeBaseSearchRequest, db: Session = Depends(database.get_db)):
	item = db.query(models.KnowledgeBase).filter(models.KnowledgeBase.id == knowledge_base_id).first()
	if item is None:
		raise HTTPException(status_code=404, detail="Knowledge base not found")
	if not item.remote_resource_id:
		raise HTTPException(status_code=400, detail="Knowledge base is not linked to a valid remote Volcengine resource_id. Reconnect this knowledge base before searching.")

	try:
		search_payload = search_knowledge(
			query=payload.query,
			resource_id=item.remote_resource_id,
			name=item.remote_collection_name,
			project=item.remote_project,
			host=item.remote_host,
			messages=payload.messages,
			limit=payload.limit,
			rewrite=payload.rewrite,
			rerank_switch=payload.rerank_switch,
			chunk_group=payload.chunk_group,
			retrieve_count=payload.retrieve_count,
		)
	except Exception as exc:
		item.sync_status = "search_error"
		item.sync_error = str(exc)
		sync_info = dict(item.sync_info or {})
		sync_info["status"] = item.sync_status
		item.sync_info = sync_info or None
		db.commit()
		raise HTTPException(status_code=502, detail=str(exc)) from exc

	item.sync_status = "connected"
	item.sync_error = None
	sync_info = dict(item.sync_info or {})
	if sync_info:
		sync_info["status"] = item.sync_status
		sync_info["request_id"] = search_payload.get("request_id")
		item.sync_info = sync_info
	db.commit()
	hits, meta = format_search_hits(
		knowledge_base_id=item.id,
		knowledge_base_name=item.local_display_name or item.name,
		search_payload=search_payload,
	)
	return schemas.KnowledgeBaseSearchResponse(
		query=payload.query,
		count=meta.get("count") or len(hits),
		rewrite_query=meta.get("rewrite_query"),
		request_id=meta.get("request_id"),
		token_usage=meta.get("token_usage"),
		result_list=[schemas.KnowledgeSearchHit(**hit) for hit in hits],
	)


@router.post("/{knowledge_base_id}/documents/import", response_model=schemas.KnowledgeDocumentImportResponse)
def import_document_to_knowledge_base(
	knowledge_base_id: str,
	payload: schemas.KnowledgeDocumentImportRequest,
	db: Session = Depends(database.get_db),
):
	item = _get_knowledge_base_or_404(knowledge_base_id, db)
	if not item.remote_resource_id:
		raise HTTPException(status_code=400, detail="Knowledge base is not linked to a valid remote Volcengine resource_id. Reconnect this knowledge base before importing documents.")

	uri = str(payload.uri or "").strip()
	if not uri:
		raise HTTPException(status_code=400, detail="Document uri is required")
	if not (uri.startswith("http://") or uri.startswith("https://") or uri.startswith("tos://")):
		raise HTTPException(
			status_code=400,
			detail="Volcengine document import currently requires a public URL or TOS URI. Local file upload still needs a staging step.",
		)

	document_name = str(payload.doc_name or "").strip() or _default_doc_name_from_uri(uri)
	document_type = str(payload.doc_type or infer_doc_type(doc_name=document_name, uri=uri) or "").strip().lower() or None
	remote_doc_id = str(payload.doc_id or f"doc_{uuid4().hex[:16]}").strip()
	if db.query(models.KnowledgeDocument).filter(models.KnowledgeDocument.id == remote_doc_id).first() is not None:
		remote_doc_id = f"doc_{uuid4().hex[:16]}"

	tags = [tag.model_dump() for tag in payload.tag_list]
	try:
		remote_result = add_document_by_uri(
			resource_id=item.remote_resource_id,
			collection_name=item.remote_collection_name,
			project=item.remote_project,
			host=item.remote_host,
			doc_id=remote_doc_id,
			doc_name=document_name,
			doc_type=document_type,
			uri=uri,
			description=payload.description or "",
			tag_list=tags,
		)
	except Exception as exc:
		item.sync_status = "document_import_error"
		item.sync_error = str(exc)
		sync_info = dict(item.sync_info or {})
		sync_info["status"] = item.sync_status
		item.sync_info = sync_info or None
		db.commit()
		raise HTTPException(status_code=502, detail=str(exc)) from exc

	document = _persist_imported_document(
		item=item,
		db=db,
		remote_result=remote_result,
		document_name=document_name,
		document_type=document_type,
		source_url=uri,
		tags=tags,
		size_bytes=0,
	)

	return schemas.KnowledgeDocumentImportResponse(
		knowledge_base_id=item.id,
		request_id=remote_result.get("request_id"),
		resource_id=remote_result.get("resource_id"),
		collection_name=remote_result.get("collection_name"),
		project=remote_result.get("project"),
		document=_serialize_document(document),
	)


@router.post("/{knowledge_base_id}/documents/upload", response_model=schemas.KnowledgeDocumentUploadResponse)
async def upload_document_to_tos_and_import(
	knowledge_base_id: str,
	file: UploadFile = File(...),
	doc_id: str | None = Form(default=None),
	doc_name: str | None = Form(default=None),
	doc_type: str | None = Form(default=None),
	description: str = Form(default=""),
	db: Session = Depends(database.get_db),
):
	item = _get_knowledge_base_or_404(knowledge_base_id, db)
	if not item.remote_resource_id:
		raise HTTPException(status_code=400, detail="Knowledge base is not linked to a valid remote Volcengine resource_id. Reconnect this knowledge base before uploading documents.")

	file_bytes = await file.read()
	if not file_bytes:
		raise HTTPException(status_code=400, detail="Uploaded file is empty")

	resolved_doc_name = str(doc_name or file.filename or "document").strip() or "document"
	resolved_doc_type = str(doc_type or infer_doc_type(doc_name=resolved_doc_name, uri=resolved_doc_name) or "").strip().lower() or None
	remote_doc_id = str(doc_id or f"doc_{uuid4().hex[:16]}").strip()

	try:
		tos_result = upload_bytes_to_tos(
			knowledge_base_id=knowledge_base_id,
			file_name=resolved_doc_name,
			content=file_bytes,
			content_type=file.content_type,
		)
		remote_result = add_document_by_uri(
			resource_id=item.remote_resource_id,
			collection_name=item.remote_collection_name,
			project=item.remote_project,
			host=item.remote_host,
			doc_id=remote_doc_id,
			doc_name=resolved_doc_name,
			doc_type=resolved_doc_type,
			uri=tos_result["uri"],
			description=description or "",
			tag_list=[],
		)
	except Exception as exc:
		item.sync_status = "document_upload_error"
		item.sync_error = str(exc)
		sync_info = dict(item.sync_info or {})
		sync_info["status"] = item.sync_status
		item.sync_info = sync_info or None
		db.commit()
		raise HTTPException(status_code=502, detail=str(exc)) from exc

	document = _persist_imported_document(
		item=item,
		db=db,
		remote_result={**remote_result, "uri": tos_result["uri"]},
		document_name=resolved_doc_name,
		document_type=resolved_doc_type,
		source_url=tos_result["uri"],
		tags=[],
		size_bytes=len(file_bytes),
	)

	return schemas.KnowledgeDocumentUploadResponse(
		knowledge_base_id=item.id,
		tos_uri=tos_result["uri"],
		bucket=tos_result["bucket"],
		key=tos_result["key"],
		request_id=remote_result.get("request_id"),
		resource_id=remote_result.get("resource_id"),
		collection_name=remote_result.get("collection_name"),
		project=remote_result.get("project"),
		document=_serialize_document(document),
	)
