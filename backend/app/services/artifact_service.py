from __future__ import annotations

from datetime import datetime
from html import escape
from io import BytesIO
from pathlib import Path
from typing import Iterable
import re
from html.parser import HTMLParser

from fastapi import HTTPException

try:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfbase.pdfmetrics import registerFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
except ImportError:  # pragma: no cover
    A4 = None
    ParagraphStyle = None
    getSampleStyleSheet = None
    UnicodeCIDFont = None
    registerFont = None
    Paragraph = None
    SimpleDocTemplate = None
    Spacer = None

try:
    from pptx import Presentation
    from pptx.util import Pt
except ImportError:  # pragma: no cover
    Presentation = None
    Pt = None

try:
    from openpyxl import Workbook, load_workbook
except ImportError:  # pragma: no cover
    Workbook = None
    load_workbook = None

try:
    from pypdf import PdfReader
except ImportError:  # pragma: no cover
    PdfReader = None


PROJECT_ROOT = Path(__file__).resolve().parents[3]
WORKSPACE_ROOT = PROJECT_ROOT / "workspace"
ARTIFACT_ROOT = PROJECT_ROOT / "workspace" / "generated"
TEXT_PREVIEW_SUFFIXES = {".txt", ".md", ".markdown", ".json", ".csv", ".html", ".htm"}
FORMAT_MIME_TYPES = {
    "html": "text/html; charset=utf-8",
    "md": "text/markdown; charset=utf-8",
    "pdf": "application/pdf",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "jpg": "image/jpeg",
    "jpeg": "image/jpeg",
}
FORMAT_EXTENSIONS = {
    "html": ".html",
    "md": ".md",
    "pdf": ".pdf",
    "pptx": ".pptx",
    "xlsx": ".xlsx",
}
FORMAT_PATTERNS = {
    "html": [r"\bhtml\b", r"html", r"\.html", r"网页", r"页面"],
    "md": [r"\bmarkdown\b", r"markdown", r"\bmd\b", r"md", r"\.md"],
    "pdf": [r"\bpdf\b", r"pdf", r"\.pdf"],
    "pptx": [r"\bpptx\b", r"pptx", r"\bppt\b", r"ppt", r"\.pptx", r"\.ppt", r"powerpoint", r"演示文稿", r"幻灯片"],
    "xlsx": [r"\bxlsx\b", r"\.xlsx", r"\bxls\b", r"\.xls", r"excel", r"spreadsheet", r"工作簿", r"电子表格", r"表格"],
}
EMPHASIS_LABEL_PATTERN = re.compile(
    r"(?m)^(?P<prefix>\s*(?:[-*]\s*)?)(?P<label>执行摘要|执行结果|关键信息|核心洞察|策略建议|目标|完成标准|结论|建议|背景|研究目标|报告框架|视觉哲学|整体结构|视觉与配色方案|布局与层次|交互与响应式设计|视觉风格特征|全球与中国市场规模与趋势|区域与品类结构|行业AI与科技应用|企业与资本动态|消费者与社会趋势|宠物食品|宠物医疗|宠物服务|宠物用品|宠物科技|市场方向|产品创新|数字化与智能化|品牌与出海策略|资本与生态布局|主色|辅色|字体|背景|增长驱动|未满足需求|AI机会)(?P<colon>[：:])"
)
TOP_LEVEL_SECTION_PATTERN = re.compile(r"^[一二三四五六七八九十]+、\s*.+$")
STANDALONE_NUMBERED_HEADING_PATTERN = re.compile(r"^\d+\.\s*[^。！？!?：:]{1,36}[：:]?$")
STANDALONE_PAREN_HEADING_PATTERN = re.compile(r"^（\d+）\s*[^。！？!?：:]{1,36}$")
STANDALONE_CHINESE_PAREN_HEADING_PATTERN = re.compile(r"^（[一二三四五六七八九十]+）\s*[^。！？!?：:]{1,40}$")
EMPHASIZED_LABEL_LINE_PATTERN = re.compile(r"^\*\*(执行摘要|执行结果|关键信息|核心洞察|策略建议|目标|完成标准|结论|建议|背景|研究目标|报告框架|下一步)[：:]\*\*$")


def detect_requested_output_formats(user_request: str) -> list[str]:
    normalized = (user_request or "").lower()
    matched_formats: list[str] = []
    for artifact_format, patterns in FORMAT_PATTERNS.items():
        if any(re.search(pattern, normalized, flags=re.IGNORECASE) for pattern in patterns):
            matched_formats.append(artifact_format)
    return matched_formats


def create_artifacts_for_response(*, user_request: str, response_text: str) -> list[dict]:
    formats = detect_requested_output_formats(user_request)
    cleaned_response = (response_text or "").strip()
    if not formats or not cleaned_response:
        return []

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    title = _derive_title(user_request, cleaned_response)
    artifact_dir = ARTIFACT_ROOT / timestamp
    artifact_dir.mkdir(parents=True, exist_ok=True)

    artifacts: list[dict] = []
    for artifact_format in formats:
        artifact_content = cleaned_response
        preserve_markdown_body = False
        if artifact_format == "html":
            embedded_html_content = extract_embedded_artifact_content(cleaned_response, "html")
            if embedded_html_content:
                artifact_content = embedded_html_content
        elif artifact_format == "md":
            embedded_markdown_content = extract_embedded_artifact_content(cleaned_response, "md")
            if embedded_markdown_content:
                artifact_content = embedded_markdown_content
                if _looks_like_jd_markdown(user_request, artifact_content):
                    artifact_content = _normalize_jd_markdown_attachment(artifact_content)
                preserve_markdown_body = True

        file_name = f"{title}_{artifact_format}{FORMAT_EXTENSIONS[artifact_format]}"
        file_path = artifact_dir / file_name

        if artifact_format == "html":
            file_path.write_text(_build_html_document(title, artifact_content), encoding="utf-8")
        elif artifact_format == "md":
            file_path.write_text(_build_markdown_document(title, artifact_content, preserve_body=preserve_markdown_body), encoding="utf-8")
        elif artifact_format == "pdf":
            _build_pdf_document(file_path, title, artifact_content)
        elif artifact_format == "pptx":
            _build_pptx_document(file_path, title, artifact_content)
        elif artifact_format == "xlsx":
            _build_xlsx_document(file_path, title, artifact_content)
        else:
            continue

        artifacts.append(_build_artifact_metadata(file_path, artifact_format, title, artifact_content))

    return artifacts


def create_artifact_from_content(*, artifact_format: str, content: str, title: str | None = None, user_request: str = "") -> dict | None:
    cleaned_content = (content or "").strip()
    normalized_format = str(artifact_format or "").lower().strip()
    if not cleaned_content or normalized_format not in FORMAT_EXTENSIONS:
        return None

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    resolved_title = title or _derive_title(user_request, cleaned_content)
    artifact_dir = ARTIFACT_ROOT / timestamp
    artifact_dir.mkdir(parents=True, exist_ok=True)

    file_name = f"{resolved_title}_{normalized_format}{FORMAT_EXTENSIONS[normalized_format]}"
    file_path = artifact_dir / file_name

    if normalized_format == "html":
        file_path.write_text(_build_html_document(resolved_title, cleaned_content), encoding="utf-8")
    elif normalized_format == "md":
        file_path.write_text(_build_markdown_document(resolved_title, cleaned_content), encoding="utf-8")
    elif normalized_format == "pdf":
        _build_pdf_document(file_path, resolved_title, cleaned_content)
    elif normalized_format == "pptx":
        _build_pptx_document(file_path, resolved_title, cleaned_content)
    elif normalized_format == "xlsx":
        _build_xlsx_document(file_path, resolved_title, cleaned_content)
    else:
        return None

    return _build_artifact_metadata(file_path, normalized_format, resolved_title, cleaned_content)


def extract_embedded_artifact_content(text: str, artifact_format: str) -> str | None:
    normalized_text = text or ""
    normalized_format = str(artifact_format or "").lower().strip()
    if not normalized_text:
        return None

    if normalized_format == "html":
        html_match = re.search(r"<!DOCTYPE\s+html[\s\S]*?</html>|<html[\s\S]*?</html>", normalized_text, flags=re.IGNORECASE)
        if html_match:
            return html_match.group(0).strip()

    if normalized_format == "md":
        markdown_match = re.search(r"```markdown\s*([\s\S]*?)```|```md\s*([\s\S]*?)```", normalized_text, flags=re.IGNORECASE)
        if markdown_match:
            return _clean_embedded_markdown_content(markdown_match.group(1) or markdown_match.group(2) or "")

        attachment_match = re.search(
            r"附件(?:内容|正文|源文档)?[：:]\s*(?:markdown|md)?\s*([\s\S]+)$",
            normalized_text,
            flags=re.IGNORECASE,
        )
        if attachment_match:
            return _clean_embedded_markdown_content(attachment_match.group(1) or "")

    return None


class _PptxHtmlTextExtractor(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts: list[str] = []
        self.skip_depth = 0

    def handle_starttag(self, tag: str, attrs) -> None:
        if tag.lower() in {"script", "style", "noscript", "svg"}:
            self.skip_depth += 1

    def handle_endtag(self, tag: str) -> None:
        if tag.lower() in {"script", "style", "noscript", "svg"} and self.skip_depth:
            self.skip_depth -= 1

    def handle_data(self, data: str) -> None:
        if self.skip_depth:
            return
        cleaned = re.sub(r"\s+", " ", data).strip()
        if cleaned:
            self.parts.append(cleaned)


def _prepare_pptx_content(body_text: str) -> str:
    embedded_html = extract_embedded_artifact_content(body_text, "html")
    if not embedded_html:
        return body_text

    parser = _PptxHtmlTextExtractor()
    parser.feed(embedded_html)
    lines: list[str] = []
    for part in parser.parts:
        if part not in lines:
            lines.append(part)
    return "\n".join(lines) or body_text


def _clean_embedded_markdown_content(content: str) -> str | None:
    cleaned = (content or "").strip()
    if not cleaned:
        return None
    cleaned = re.sub(r"(?is)^\s*(?:markdown|md)\b\s*", "", cleaned).strip()
    cleaned = re.sub(r"(?is)\s*```\s*$", "", cleaned).strip()
    return cleaned or None


JD_MARKDOWN_SECTIONS = ("岗位名称", "所属部门", "岗位概述", "岗位职责", "任职要求", "加分项", "核心能力", "合规提示")
JD_LIST_SECTIONS = {"岗位职责", "任职要求"}


def _looks_like_jd_markdown(user_request: str, content: str) -> bool:
    text = f"{user_request}\n{content}".lower()
    return ("jd" in text or "岗位" in text) and ("岗位职责" in text or "任职要求" in text)


def _normalize_jd_markdown_attachment(content: str) -> str:
    lines = [line.strip() for line in (content or "").splitlines()]
    normalized_lines: list[str] = []
    current_section = ""
    list_index = 1

    def append_line(line: str = "") -> None:
        if line or (normalized_lines and normalized_lines[-1]):
            normalized_lines.append(line)

    for line in lines:
        if not line:
            append_line("")
            continue

        if not normalized_lines and not line.startswith("#"):
            append_line(f"# {line}")
            append_line("")
            continue

        section_match = next(
            (
                section
                for section in JD_MARKDOWN_SECTIONS
                if line == section or line.startswith(f"{section} ") or line.startswith(f"## {section}")
            ),
            "",
        )
        if section_match:
            current_section = section_match
            list_index = 1
            value = re.sub(rf"^#*\s*{re.escape(section_match)}\s*", "", line).strip()
            if normalized_lines and normalized_lines[-1]:
                append_line("")
            append_line(f"## {section_match}")
            if value:
                append_line(value)
            continue

        if current_section in JD_LIST_SECTIONS:
            list_item_match = re.match(r"^(?:#{3,6}\s*|\d+[.、]\s*|[-*]\s*)(.+)$", line)
            if list_item_match:
                append_line(f"{list_index}. {list_item_match.group(1).strip()}")
                list_index += 1
                continue

        append_line(line)

    return "\n".join(normalized_lines).strip() + "\n"


def should_create_step_artifacts(*, user_request: str, step_content: str) -> bool:
    formats = detect_requested_output_formats(user_request)
    if not formats:
        return False

    normalized_step = (step_content or "").lower()
    if any(token in normalized_step for token in ["文件", "导出", "输出", "保存", "交付"]):
        return True

    for artifact_format in formats:
        if any(re.search(pattern, normalized_step, flags=re.IGNORECASE) for pattern in FORMAT_PATTERNS[artifact_format]):
            return True

    return False


def create_step_result_md_artifact(
    *,
    step_index: int,
    step_name: str,
    result_text: str,
    detailed_text: str = "",
    user_request: str = "",
) -> dict | None:
    """Save a completed step's result as a Markdown file and return artifact metadata."""
    cleaned = format_summary_for_display(_extract_summary_content((detailed_text or result_text or "").strip()))
    if not cleaned:
        return None

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    artifact_dir = ARTIFACT_ROOT / timestamp
    artifact_dir.mkdir(parents=True, exist_ok=True)

    safe_name = re.sub(r"[^\w\u4e00-\u9fff]", "_", (step_name or "")[:40]).strip("_")
    safe_name = re.sub(r"_+", "_", safe_name) or f"step_{step_index}"
    file_name = f"Step{step_index}_{safe_name}.md"
    file_path = artifact_dir / file_name

    md_content = f"# {step_name or f'步骤 {step_index}'}\n\n{cleaned}\n"
    file_path.write_text(md_content, encoding="utf-8")

    preview = (result_text or cleaned)[:4000]
    return _build_artifact_metadata(file_path, "md", step_name or f"步骤 {step_index}", preview)


def aggregate_step_results_for_final(
    *,
    step_outputs: dict,
    current_step_result: str = "",
    user_request: str,
) -> list[dict]:
    """Combine all step results into the final deliverable file(s)."""
    ordered_entries = sorted(
        ((k, v) for k, v in (step_outputs or {}).items() if isinstance(v, dict)),
        key=lambda item: _extract_step_sort_key(item[0]),
    )

    rendered_artifacts = render_requested_deliverables(
        step_outputs=step_outputs,
        current_step_result=current_step_result,
        user_request=user_request,
    )
    if rendered_artifacts:
        return rendered_artifacts

    requested_formats = detect_requested_output_formats(user_request)
    source_texts: list[str] = []
    current = _extract_summary_content(current_step_result)
    if current:
        source_texts.append(current)
    source_texts.extend(
        _extract_summary_content(entry.get("result") or entry.get("summary") or "")
        for _, entry in reversed(ordered_entries)
        if _extract_summary_content(entry.get("result") or entry.get("summary") or "")
    )

    preferred_artifacts: list[dict] = []
    for artifact_format in requested_formats:
        for source_text in source_texts:
            embedded_content = extract_embedded_artifact_content(source_text, artifact_format)
            if not embedded_content:
                continue

            artifact = create_artifact_from_content(
                artifact_format=artifact_format,
                content=embedded_content,
                user_request=user_request,
            )
            if artifact:
                preferred_artifacts.append(artifact)
                break

    if preferred_artifacts:
        return preferred_artifacts

    parts: list[str] = []
    for _step_id, entry in ordered_entries:
        text = _extract_full_step_text(entry)
        if text:
            parts.append(text)

    if current:
        parts.append(current)

    combined_text = "\n\n".join(parts)
    if not combined_text:
        return []

    return create_artifacts_for_response(
        user_request=user_request,
        response_text=combined_text,
    )


def build_step_context_summary_for_prompt(
    step_outputs: dict,
    *,
    max_steps: int = 6,
    max_chars_per_step: int = 2200,
    max_total_chars: int = 12000,
) -> str:
    ordered_entries = sorted(
        ((k, v) for k, v in (step_outputs or {}).items() if isinstance(v, dict)),
        key=lambda item: _extract_step_sort_key(item[0]),
    )
    if not ordered_entries:
        return ""

    selected_entries = ordered_entries[-max_steps:]
    omitted_count = max(0, len(ordered_entries) - len(selected_entries))
    remaining_chars = max_total_chars
    sections: list[str] = []

    if omitted_count:
        sections.append(f"已省略更早的 {omitted_count} 个步骤，仅保留最近且可复用的关键结果。")
        remaining_chars -= len(sections[0])

    for step_id, entry in selected_entries:
        title = str(entry.get("title") or "").strip() or f"步骤 {step_id}"
        summary = _normalize_prompt_context_text(_extract_full_step_text(entry))
        if not summary:
            continue

        artifact_names = [
            str(artifact.get("name") or artifact.get("title") or "").strip()
            for artifact in normalize_artifacts(entry.get("artifacts") or [])
            if str(artifact.get("name") or artifact.get("title") or "").strip()
        ]

        available_chars = min(max_chars_per_step, max(0, remaining_chars - 120))
        if available_chars <= 0:
            break

        if len(summary) > available_chars:
            summary = summary[: max(available_chars - 3, 0)].rstrip()
            if summary:
                summary += "..."

        block_lines = [
            f"[{step_id}] {title}",
            f"摘要：{summary}",
        ]
        if artifact_names:
            block_lines.append(f"相关产物：{', '.join(dict.fromkeys(artifact_names[:4]))}")

        block = "\n".join(block_lines)
        if len(block) > remaining_chars and sections:
            break

        sections.append(block)
        remaining_chars -= len(block)
        if remaining_chars <= 0:
            break

    return "\n\n".join(section for section in sections if section).strip()


def _extract_full_step_text(entry: dict) -> str:
    if not isinstance(entry, dict):
        return ""

    for artifact in normalize_artifacts(entry.get("artifacts") or []):
        relative_path = str(artifact.get("relative_path") or "").strip()
        artifact_format = str(artifact.get("format") or "").lower()
        if not relative_path or artifact_format not in {"md", "markdown", "txt"}:
            continue

        file_path = resolve_workspace_file_path(relative_path)
        if file_path.exists() and file_path.is_file():
            try:
                return _extract_summary_content(file_path.read_text(encoding="utf-8"))
            except OSError:
                continue

    return _extract_summary_content(entry.get("result") or entry.get("summary") or "")


def _normalize_prompt_context_text(text: str) -> str:
    normalized = _extract_summary_content(text)
    normalized = re.sub(r"\n{3,}", "\n\n", normalized)
    normalized = re.sub(r"[ \t]+", " ", normalized)
    return normalized.strip()


def _extract_step_sort_key(step_id: str) -> int:
    match = re.search(r"(\d+)", str(step_id))
    return int(match.group(1)) if match else 0


def normalize_artifacts(raw_artifacts: Iterable[dict] | None) -> list[dict]:
    normalized: list[dict] = []
    for raw_artifact in raw_artifacts or []:
        artifact = raw_artifact.model_dump() if hasattr(raw_artifact, "model_dump") else raw_artifact
        relative_path = str(artifact.get("relative_path") or "").replace("\\", "/").strip()
        if not relative_path:
            continue

        file_path = resolve_workspace_file_path(relative_path)
        if not file_path.exists():
            continue

        artifact_format = str(artifact.get("format") or file_path.suffix.lstrip(".") or "file").lower()
        if artifact_format not in FORMAT_MIME_TYPES:
            continue
        normalized.append({
            "id": artifact.get("id") or relative_path,
            "name": artifact.get("name") or file_path.name,
            "title": artifact.get("title") or file_path.stem,
            "relative_path": relative_path,
            "format": artifact_format,
            "mime_type": artifact.get("mime_type") or FORMAT_MIME_TYPES.get(artifact_format, "application/octet-stream"),
            "size_bytes": int(artifact.get("size_bytes") or file_path.stat().st_size),
            "created_at": artifact.get("created_at") or datetime.fromtimestamp(file_path.stat().st_mtime).isoformat(),
            "preview_text": artifact.get("preview_text") or "",
        })
    return normalized


def build_artifact_metadata_from_file_path(file_path: str | Path, *, title: str | None = None, preview_text: str = "") -> dict | None:
    candidate = Path(file_path)
    if not candidate.is_absolute():
        candidate = (WORKSPACE_ROOT / candidate).resolve()
    else:
        candidate = candidate.resolve()

    allowed_root = WORKSPACE_ROOT.resolve()
    if not str(candidate).startswith(str(allowed_root)):
        return None
    if not candidate.exists() or not candidate.is_file():
        return None

    artifact_format = candidate.suffix.lstrip(".").lower() or "file"
    if artifact_format not in FORMAT_MIME_TYPES:
        return None
    resolved_title = title or candidate.stem
    return _build_artifact_metadata(candidate, artifact_format, resolved_title, preview_text)


def is_valid_native_pptx(file_path: str | Path) -> bool:
    if Presentation is None:
        return False
    candidate = Path(file_path)
    if not candidate.exists() or not candidate.is_file() or candidate.suffix.lower() != ".pptx":
        return False
    try:
        presentation = Presentation(str(candidate))
    except Exception:
        return False
    return len(presentation.slides) > 0


def is_valid_native_xlsx(file_path: str | Path) -> bool:
    if load_workbook is None:
        return False
    candidate = Path(file_path)
    if not candidate.is_file() or candidate.suffix.lower() != ".xlsx":
        return False
    try:
        workbook = load_workbook(candidate, read_only=True, data_only=False)
        valid = len(workbook.sheetnames) > 0
        workbook.close()
        return valid
    except Exception:
        return False


def is_valid_native_pdf(file_path: str | Path) -> bool:
    candidate = Path(file_path)
    if not candidate.is_file() or candidate.suffix.lower() != ".pdf":
        return False
    try:
        if candidate.read_bytes()[:5] != b"%PDF-":
            return False
        if PdfReader is None:
            return candidate.stat().st_size > 100
        return len(PdfReader(str(candidate)).pages) > 0
    except Exception:
        return False


def resolve_workspace_file_path(relative_path: str) -> Path:
    normalized = str(relative_path or "").replace("\\", "/").lstrip("/")
    candidate = (PROJECT_ROOT / normalized).resolve()
    allowed_root = WORKSPACE_ROOT.resolve()
    if not str(candidate).startswith(str(allowed_root)):
        raise HTTPException(status_code=400, detail="Workspace file path is not allowed")
    return candidate


def read_workspace_text_file(relative_path: str) -> dict:
    file_path = resolve_workspace_file_path(relative_path)
    if not file_path.exists() or not file_path.is_file():
        raise HTTPException(status_code=404, detail="Workspace file not found")
    return {
        "name": file_path.name,
        "relative_path": str(file_path.relative_to(PROJECT_ROOT)).replace("\\", "/"),
        "mime_type": guess_mime_type(file_path),
        "content": file_path.read_text(encoding="utf-8"),
    }


def is_text_previewable(relative_path: str) -> bool:
    return resolve_workspace_file_path(relative_path).suffix.lower() in TEXT_PREVIEW_SUFFIXES


def guess_mime_type(file_path: Path) -> str:
    suffix = file_path.suffix.lower()
    if suffix in {".html", ".htm"}:
        return FORMAT_MIME_TYPES["html"]
    if suffix in {".md", ".markdown"}:
        return FORMAT_MIME_TYPES["md"]
    if suffix == ".pdf":
        return FORMAT_MIME_TYPES["pdf"]
    if suffix == ".pptx":
        return FORMAT_MIME_TYPES["pptx"]
    if suffix == ".xlsx":
        return FORMAT_MIME_TYPES["xlsx"]
    if suffix in {".jpg", ".jpeg"}:
        return FORMAT_MIME_TYPES[suffix.lstrip(".")]
    return "text/plain; charset=utf-8"


def _build_xlsx_document(file_path: Path, title: str, body_text: str) -> None:
    if Workbook is None:
        raise HTTPException(status_code=500, detail="XLSX generation dependency is not installed")

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = _sanitize_excel_sheet_title(title)

    rows = _parse_text_to_excel_rows(body_text)
    for row_index, row in enumerate(rows, start=1):
        for column_index, value in enumerate(row, start=1):
            worksheet.cell(row=row_index, column=column_index, value=value)

    for column_cells in worksheet.columns:
        values = ["" if cell.value is None else str(cell.value) for cell in column_cells]
        max_length = max((len(value) for value in values), default=0)
        worksheet.column_dimensions[column_cells[0].column_letter].width = min(max(max_length + 2, 10), 48)

    worksheet.freeze_panes = "A2"
    workbook.save(file_path)


def _sanitize_excel_sheet_title(title: str) -> str:
    cleaned = re.sub(r'[\\/*?:\[\]]', "_", (title or "Sheet1")).strip()
    return (cleaned[:31] or "Sheet1")


def _parse_text_to_excel_rows(body_text: str) -> list[list[str]]:
    normalized = (body_text or "").replace("\r\n", "\n").strip()
    if not normalized:
        return [[""]]

    markdown_rows = _parse_markdown_table_rows(normalized)
    if markdown_rows:
        return markdown_rows

    csv_rows = _parse_delimited_rows(normalized, ",")
    if csv_rows:
        return csv_rows

    tsv_rows = _parse_delimited_rows(normalized, "\t")
    if tsv_rows:
        return tsv_rows

    lines = [line.strip() for line in normalized.split("\n") if line.strip()]
    return [[line] for line in lines] or [[normalized]]


def _parse_markdown_table_rows(text: str) -> list[list[str]]:
    lines = [line.strip() for line in text.split("\n") if line.strip()]
    pipe_lines = [line for line in lines if "|" in line]
    if len(pipe_lines) < 2:
        return []

    rows: list[list[str]] = []
    for line in pipe_lines:
        cells = [cell.strip() for cell in line.strip("|").split("|")]
        if not any(cells):
            continue
        if all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells):
            continue
        rows.append(cells)
    return rows


def _parse_delimited_rows(text: str, delimiter: str) -> list[list[str]]:
    lines = [line for line in text.split("\n") if line.strip()]
    if len(lines) < 2:
        return []
    if not all(delimiter in line for line in lines[: min(len(lines), 5)]):
        return []
    return [[cell.strip() for cell in line.split(delimiter)] for line in lines]


def _build_artifact_metadata(file_path: Path, artifact_format: str, title: str, preview_text: str) -> dict:
    if artifact_format not in FORMAT_MIME_TYPES:
        raise ValueError(f"Unsupported artifact format: {artifact_format}")
    relative_path = str(file_path.relative_to(PROJECT_ROOT)).replace("\\", "/")
    return {
        "id": relative_path,
        "name": file_path.name,
        "title": title,
        "relative_path": relative_path,
        "format": artifact_format,
        "mime_type": FORMAT_MIME_TYPES[artifact_format],
        "size_bytes": file_path.stat().st_size,
        "created_at": datetime.now().isoformat(),
        "preview_text": preview_text[:4000],
    }


def _derive_title(user_request: str, response_text: str) -> str:
    request_text = _extract_summary_content(user_request)
    response_text = _strip_inline_source_titles(_extract_summary_content(response_text))

    source = _match_title_candidate(request_text)
    if not source:
        source = _match_title_candidate(response_text)
    if not source:
        source = (request_text or response_text or "deliverable").strip()

    source = re.sub(r"\s+", "_", source)
    source = re.sub(r"[^0-9A-Za-z_\-\u4e00-\u9fff]", "", source)
    source = source.strip("_")
    return (source or "deliverable")[:32]


def _extract_summary_content(text: str) -> str:
    normalized = (text or "").strip()
    if not normalized:
        return ""

    summary_match = re.search(r"##\s*简要执行结果\s*(.*)$", normalized, flags=re.IGNORECASE | re.DOTALL)
    if summary_match:
        normalized = summary_match.group(1).strip()

    normalized = re.sub(r"\n##\s*完整执行记录[\s\S]*$", "", normalized, flags=re.IGNORECASE).strip()
    normalized = re.sub(r"\n###\s*消息\s+\d+\s+·[\s\S]*$", "", normalized, flags=re.IGNORECASE).strip()
    normalized = re.sub(r"\n###\s*工具调用\s+\d+(?:\.\d+)?\s+·[\s\S]*$", "", normalized, flags=re.IGNORECASE).strip()
    normalized = re.sub(r"(?mi)^\s*相关产物：.*$", "", normalized)
    normalized = re.sub(r"(?mi)^\s*如果后续步骤需要依赖本结果.*$", "", normalized)
    normalized = re.sub(r'(?mi)^\s*-\s*Report\s+""\s+when\s+this\s+step\s+is\s+completed\.\s*$', "", normalized)
    normalized = re.sub(r"(?mi)^\s*###\s*消息\s+\d+\s+·.*$", "", normalized)
    normalized = re.sub(r"(?mi)^\s*###\s*工具调用\s+\d+(?:\.\d+)?\s+·.*$", "", normalized)
    normalized = re.sub(r"(?s)\n(?:下一步：|下一步:)\s*.*$", "", normalized)
    normalized = re.sub(r"(?m)^\s*执行结果：\s*", "", normalized)
    normalized = re.sub(r"(?m)^\s*关键信息：\s*$", "", normalized)
    normalized = _strip_embedded_html_source(normalized)
    normalized = re.sub(r"\n{3,}", "\n\n", normalized)

    return normalized.strip()


def format_summary_for_display(text: str) -> str:
    normalized = _extract_summary_content(text)
    if not normalized:
        return ""

    normalized = EMPHASIS_LABEL_PATTERN.sub(
        lambda match: f"{match.group('prefix')}**{match.group('label')}{match.group('colon')}**",
        normalized,
    )
    paragraphs = [segment.strip() for segment in re.split(r"\n\s*\n", normalized) if segment.strip()]
    looks_like_structured_summary = len(paragraphs) > 1 or bool(re.search(r"(?m)^(?:[-*]|\d+[\.、])\s+", normalized))
    if paragraphs:
        first_paragraph = paragraphs[0]
        if (
            looks_like_structured_summary
            and
            not first_paragraph.startswith(("#", "**"))
            and not re.match(r"^(?:[-*]|\d+[\.、])\s+", first_paragraph)
            and len(first_paragraph) <= 120
        ):
            paragraphs[0] = f"**执行结果：**\n{first_paragraph}"
            normalized = "\n\n".join(paragraphs)
    normalized = re.sub(r"\*\*[ \t]+", "**", normalized)
    normalized = re.sub(r"[ \t]+\*\*", "**", normalized)
    normalized = _normalize_markdown_layout(normalized)
    return normalized.strip()


def _normalize_markdown_layout(text: str) -> str:
    lines = (text or "").splitlines()
    rendered: list[str] = []
    top_level_index = 0
    sub_level_index = 0

    def ensure_blank_line() -> None:
        if rendered and rendered[-1] != "":
            rendered.append("")

    def strip_section_prefix(line: str) -> str:
        normalized_line = re.sub(r"^[一二三四五六七八九十]+、\s*", "", line)
        normalized_line = re.sub(r"^（[一二三四五六七八九十\d]+）\s*", "", normalized_line)
        normalized_line = re.sub(r"^\d+\.\s*", "", normalized_line)
        return normalized_line.rstrip("：:").strip()

    for raw_line in lines:
        line = raw_line.strip()
        if not line:
            ensure_blank_line()
            continue

        if EMPHASIZED_LABEL_LINE_PATTERN.match(line):
            ensure_blank_line()
            rendered.append(line)
            rendered.append("")
            continue

        if TOP_LEVEL_SECTION_PATTERN.match(line):
            top_level_index += 1
            sub_level_index = 0
            ensure_blank_line()
            rendered.append(f"## {top_level_index}. {strip_section_prefix(line)}")
            rendered.append("")
            continue

        if STANDALONE_PAREN_HEADING_PATTERN.match(line):
            sub_level_index += 1
            ensure_blank_line()
            if top_level_index > 0:
                rendered.append(f"### {top_level_index}.{sub_level_index} {strip_section_prefix(line)}")
            else:
                rendered.append(f"### {strip_section_prefix(line)}")
            rendered.append("")
            continue

        if STANDALONE_CHINESE_PAREN_HEADING_PATTERN.match(line):
            sub_level_index += 1
            ensure_blank_line()
            if top_level_index > 0:
                rendered.append(f"### {top_level_index}.{sub_level_index} {strip_section_prefix(line)}")
            else:
                rendered.append(f"### {strip_section_prefix(line)}")
            rendered.append("")
            continue

        if STANDALONE_NUMBERED_HEADING_PATTERN.match(line):
            sub_level_index += 1
            ensure_blank_line()
            if top_level_index > 0:
                rendered.append(f"### {top_level_index}.{sub_level_index} {strip_section_prefix(line)}")
            else:
                rendered.append(f"### {strip_section_prefix(line)}")
            rendered.append("")
            continue

        parenthesized_number_match = re.match(r"^（(?P<num>\d+)）\s*(?P<content>.+)$", line)
        if parenthesized_number_match:
            if rendered:
                previous_line = rendered[-1]
                if previous_line and not previous_line.startswith("#") and not re.match(r"^(?:[-*]|\d+\.)\s+", previous_line):
                    rendered.append("")
            rendered.append(f"{parenthesized_number_match.group('num')}. {parenthesized_number_match.group('content').strip()}")
            continue

        if re.match(r"^(?:[-*]|\d+\.)\s+", line) and rendered and rendered[-1] not in {"", "-", "*"}:
            previous_line = rendered[-1]
            if not previous_line.startswith("#") and not re.match(r"^(?:[-*]|\d+\.)\s+", previous_line):
                rendered.append("")

        rendered.append(line)

    normalized = "\n".join(rendered)
    normalized = re.sub(r"\n{3,}", "\n\n", normalized)
    return normalized.strip()


def _render_inline_markdown_to_html(text: str) -> str:
    normalized = text or ""
    parts = re.split(r"(\*\*[^*]+\*\*)", normalized)
    rendered: list[str] = []
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            rendered.append(f"<strong>{escape(part[2:-2])}</strong>")
        else:
            rendered.append(escape(part))
    return "".join(rendered)


def _strip_embedded_html_source(text: str) -> str:
    lines = []
    for raw_line in (text or "").splitlines():
        line = raw_line.strip()
        if not line:
            lines.append(raw_line)
            continue

        if any(marker in line for marker in [
            "HTML报告内容",
            "【HTML报告结构】",
            "HTML报告示例结构",
            "HTML结构化摘要",
            "HTML主体示例",
        ]):
            continue

        if re.search(r"(?i)(&lt;|<)/?(html|head|body|section|header|footer|style|meta|title|nav|main|article)\b", line):
            continue

        if re.search(r"(?i)^(body|header|section|footer|nav|main|article|h1|h2|h3|p|ul|ol|li|table|thead|tbody|tr|td|th|\.highlight)\s*\{.*\}\s*$", line):
            continue

        lines.append(raw_line)

    normalized = "\n".join(lines)
    normalized = re.sub(r"\n{3,}", "\n\n", normalized)
    return normalized.strip()


def _strip_inline_source_titles(text: str) -> str:
    normalized = text or ""
    return re.sub(r"（来源：[^）]*《[^》]+》[^）]*）", "", normalized)


def _match_title_candidate(text: str) -> str:
    normalized = (text or "").strip()
    if not normalized:
        return ""

    matched_book_title = re.search(r"《([^》]{2,64})》", normalized)
    if matched_book_title:
        return matched_book_title.group(1)

    matched_heading = re.search(r"([\u4e00-\u9fffA-Za-z0-9（）()\-\s]{4,64}(报告|方案|汇报|分析|总结))", normalized)
    if matched_heading:
        return matched_heading.group(1)

    return ""


def _build_html_document(title: str, body_text: str) -> str:
    if "<html" in body_text.lower():
        return body_text

    formatted_body = format_summary_for_display(body_text)
    paragraphs = [segment.strip() for segment in re.split(r"\n\s*\n", formatted_body) if segment.strip()]
    body = "\n".join(f"<p>{_render_inline_markdown_to_html(paragraph).replace(chr(10), '<br/>')}</p>" for paragraph in paragraphs)
    return f"""<!DOCTYPE html>
<html lang=\"zh-CN\">
<head>
  <meta charset=\"UTF-8\" />
  <meta name=\"viewport\" content=\"width=device-width, initial-scale=1.0\" />
  <title>{escape(title)}</title>
  <style>
    body {{ margin: 0; padding: 36px 20px; background: linear-gradient(180deg, #f7f2e9 0%, #efe6d6 100%); color: #1f2937; font-family: Georgia, "Times New Roman", serif; }}
    .shell {{ max-width: 960px; margin: 0 auto; background: rgba(255, 252, 247, 0.92); border: 1px solid rgba(148, 163, 184, 0.22); border-radius: 28px; box-shadow: 0 24px 80px rgba(15, 23, 42, 0.12); overflow: hidden; }}
    .hero {{ padding: 44px 52px 24px; background: linear-gradient(135deg, rgba(120, 53, 15, 0.92), rgba(180, 83, 9, 0.8)); color: white; }}
    .hero h1 {{ margin: 0; font-size: 38px; line-height: 1.1; }}
    .hero p {{ margin: 12px 0 0; color: rgba(255,255,255,0.84); line-height: 1.7; }}
    .content {{ padding: 34px 52px 52px; }}
    p {{ margin: 0 0 18px; font-size: 17px; line-height: 1.9; }}
  </style>
</head>
<body>
  <main class=\"shell\">
    <section class=\"hero\">
      <h1>{escape(title)}</h1>
      <p>根据本轮对话自动生成的 HTML 交付文件。</p>
    </section>
    <section class=\"content\">{body}</section>
  </main>
</body>
</html>
"""


def build_render_context(
    *,
    step_outputs: dict,
    current_step_result: str = "",
    user_request: str,
) -> dict:
    ordered_entries = sorted(
        ((k, v) for k, v in (step_outputs or {}).items() if isinstance(v, dict)),
        key=lambda item: _extract_step_sort_key(item[0]),
    )
    step_summaries = [_extract_step_summary_entry(step_id, entry) for step_id, entry in ordered_entries]
    current_summary = _extract_summary_content(current_step_result)
    if current_summary:
        step_summaries.append({
            "step_id": "current",
            "title": "最终交付",
            "content": current_summary,
        })

    title = _derive_title(user_request, "\n\n".join(item.get("content") or "" for item in step_summaries))
    visual_spec = _extract_visual_spec(step_summaries[0].get("content", "") if step_summaries else "")
    intro_text = _extract_intro_text(step_summaries)
    requested_formats = detect_requested_output_formats(user_request)

    return {
        "title": title,
        "user_request": user_request,
        "requested_formats": requested_formats,
        "step_summaries": step_summaries,
        "visual_spec": visual_spec,
        "intro_text": intro_text,
    }


def render_requested_deliverables(
    *,
    step_outputs: dict,
    current_step_result: str = "",
    user_request: str,
) -> list[dict]:
    context = build_render_context(
        step_outputs=step_outputs,
        current_step_result=current_step_result,
        user_request=user_request,
    )
    requested_formats = context.get("requested_formats") or []
    if not requested_formats:
        return []

    step_summaries = context.get("step_summaries") or []
    if not step_summaries:
        return []

    title = context["title"]
    artifacts: list[dict] = []
    for artifact_format in requested_formats:
        if artifact_format == "html":
            html = _build_synthesized_html_document(
                title=title,
                step_summaries=step_summaries,
                visual_spec=context.get("visual_spec") or {},
            )
            artifact = create_artifact_from_content(
                artifact_format="html",
                content=html,
                title=title,
                user_request=user_request,
            )
        else:
            artifact = create_artifact_from_content(
                artifact_format=artifact_format,
                content=_build_structured_delivery_text(context),
                title=title,
                user_request=user_request,
            )

        if artifact:
            artifacts.append(artifact)

    return artifacts


def _extract_step_summary_entry(step_id: str, entry: dict) -> dict:
    title = ""
    artifacts = normalize_artifacts(entry.get("artifacts") or []) if isinstance(entry, dict) else []
    if artifacts:
        title = str(artifacts[0].get("title") or "").strip()
    if not title and isinstance(entry, dict):
        title = str(entry.get("title") or "").strip()
    if not title:
        title = f"步骤 {step_id}"

    return {
        "step_id": step_id,
        "title": title,
        "content": _extract_full_step_text(entry),
    }


def _extract_visual_spec(text: str) -> dict:
    normalized = text or ""
    title_match = re.search(r"中央(?:居中)?标题[“\"]([^”\"]{4,64})[”\"]", normalized)
    primary_match = re.search(r"主色[：:]\s*[^#\n]*(#(?:[0-9A-Fa-f]{6}))", normalized)
    accent_match = re.search(r"辅色[：:]\s*[^#\n]*(#(?:[0-9A-Fa-f]{6}))", normalized)
    background_match = re.search(r"背景[：:]\s*[^#\n]*(#(?:[0-9A-Fa-f]{6}))", normalized)

    return {
    "cover_title": title_match.group(1).strip() if title_match else "",
    "primary": primary_match.group(1) if primary_match else "#1E3A8A",
    "accent": accent_match.group(1) if accent_match else "#F59E0B",
    "background": background_match.group(1) if background_match else "#F9FAFB",
    }


def _build_synthesized_html_document(*, title: str, step_summaries: list[dict], visual_spec: dict) -> str:
    primary = visual_spec.get("primary") or "#1E3A8A"
    accent = visual_spec.get("accent") or "#F59E0B"
    background = visual_spec.get("background") or "#F9FAFB"
    cover_title = visual_spec.get("cover_title") or title.replace("_", " ")

    intro_text = _extract_intro_text(step_summaries)
    sections = []
    for item in step_summaries:
        content = _extract_summary_content(item.get("content", ""))
        if not content:
            continue
        section_title = _normalize_section_title(item.get("title") or "")
        sections.append(
            f"""
            <section class=\"report-section\"> 
                <div class=\"section-head\">
                    <span class=\"eyebrow\">{escape(item.get('step_id', ''))}</span>
                    <h2>{escape(section_title)}</h2>
                </div>
                {_render_summary_html(content)}
            </section>
"""
        )

    return f"""<!DOCTYPE html>
<html lang=\"zh-CN\">
<head>
    <meta charset=\"UTF-8\" />
    <meta name=\"viewport\" content=\"width=device-width, initial-scale=1.0\" />
    <title>{escape(title.replace('_', ' '))}</title>
    <style>
        :root {{
            --primary: {primary};
            --accent: {accent};
            --bg: {background};
            --ink: #14213d;
            --muted: #5c677d;
            --panel: rgba(255,255,255,0.88);
            --line: rgba(20, 33, 61, 0.12);
        }}
        * {{ box-sizing: border-box; }}
        body {{
            margin: 0;
            font-family: "Segoe UI", "PingFang SC", "Hiragino Sans GB", "Microsoft YaHei", sans-serif;
            color: var(--ink);
            background:
                radial-gradient(circle at top left, color-mix(in srgb, var(--accent) 20%, white) 0%, transparent 28%),
                radial-gradient(circle at top right, color-mix(in srgb, var(--primary) 18%, white) 0%, transparent 24%),
                linear-gradient(180deg, var(--bg) 0%, #eef2f7 100%);
        }}
        .shell {{ max-width: 1180px; margin: 0 auto; padding: 32px 20px 64px; }}
        .hero {{
            position: relative;
            overflow: hidden;
            border-radius: 32px;
            padding: 56px 56px 48px;
            background: linear-gradient(135deg, var(--primary) 0%, color-mix(in srgb, var(--primary) 72%, black) 100%);
            color: white;
            box-shadow: 0 30px 80px rgba(16, 24, 40, 0.22);
        }}
        .hero::after {{
            content: "";
            position: absolute;
            inset: auto -8% -28% auto;
            width: 420px;
            height: 420px;
            border-radius: 50%;
            background: radial-gradient(circle, color-mix(in srgb, var(--accent) 34%, white) 0%, transparent 68%);
            opacity: 0.45;
        }}
        .hero h1 {{ margin: 0; max-width: 760px; font-size: clamp(34px, 5vw, 56px); line-height: 1.05; letter-spacing: -0.03em; }}
        .hero p {{ margin: 18px 0 0; max-width: 720px; font-size: 17px; line-height: 1.8; color: rgba(255,255,255,0.82); }}
        .hero-meta {{ margin-top: 28px; display: flex; gap: 12px; flex-wrap: wrap; }}
        .hero-meta span {{ padding: 10px 14px; border-radius: 999px; background: rgba(255,255,255,0.12); border: 1px solid rgba(255,255,255,0.14); font-size: 13px; }}
        .grid {{ display: grid; grid-template-columns: 280px minmax(0, 1fr); gap: 24px; margin-top: 24px; align-items: start; }}
        .toc, .content {{ background: var(--panel); border: 1px solid var(--line); border-radius: 28px; box-shadow: 0 20px 60px rgba(15, 23, 42, 0.08); backdrop-filter: blur(10px); }}
        .toc {{ position: sticky; top: 20px; padding: 24px; }}
        .toc h2 {{ margin: 0 0 16px; font-size: 14px; letter-spacing: 0.18em; text-transform: uppercase; color: var(--muted); }}
        .toc ol {{ margin: 0; padding-left: 18px; }}
        .toc li {{ margin: 0 0 12px; color: var(--ink); line-height: 1.5; }}
        .content {{ padding: 30px; }}
        .summary-card {{ padding: 24px; border-radius: 24px; background: linear-gradient(180deg, rgba(255,255,255,0.96), rgba(255,255,255,0.76)); border: 1px solid var(--line); margin-bottom: 22px; }}
        .summary-card h2 {{ margin: 0 0 12px; font-size: 24px; }}
        .summary-card p {{ margin: 0; line-height: 1.9; color: var(--muted); }}
        .report-section {{ padding: 26px 0; border-top: 1px solid var(--line); }}
        .report-section:first-of-type {{ border-top: 0; padding-top: 0; }}
        .section-head {{ margin-bottom: 16px; }}
        .section-head h2 {{ margin: 6px 0 0; font-size: 30px; line-height: 1.15; }}
        .eyebrow {{ font-size: 12px; letter-spacing: 0.18em; text-transform: uppercase; color: var(--accent); font-weight: 700; }}
        .rich-text p {{ margin: 0 0 14px; line-height: 1.9; color: #314158; }}
        .rich-text ul {{ margin: 0 0 14px 0; padding-left: 22px; color: #314158; }}
        .rich-text li {{ margin-bottom: 10px; line-height: 1.85; }}
        .rich-text .callout {{ padding: 16px 18px; border-left: 4px solid var(--accent); background: color-mix(in srgb, var(--accent) 10%, white); border-radius: 0 18px 18px 0; margin: 16px 0; }}
        @media (max-width: 900px) {{
            .grid {{ grid-template-columns: 1fr; }}
            .toc {{ position: static; }}
            .hero {{ padding: 36px 28px; border-radius: 24px; }}
            .content {{ padding: 22px; }}
            .section-head h2 {{ font-size: 24px; }}
        }}
    </style>
</head>
<body>
    <main class=\"shell\">
        <section class=\"hero\">
            <h1>{escape(cover_title)}</h1>
            <p>{escape(intro_text)}</p>
            <div class=\"hero-meta\">
                <span>主题聚合交付</span>
                <span>基于步骤风格与研究内容重组</span>
                <span>输出格式：HTML</span>
            </div>
        </section>
        <div class=\"grid\">
            <aside class=\"toc\">
                <h2>Contents</h2>
                <ol>
                    {''.join(f'<li>{escape(_normalize_section_title(item.get("title") or ""))}</li>' for item in step_summaries if item.get('content'))}
                </ol>
            </aside>
            <section class=\"content\">
                <div class=\"summary-card\">
                    <h2>执行摘要</h2>
                    <p>{escape(intro_text)}</p>
                </div>
                {''.join(sections)}
            </section>
        </div>
    </main>
</body>
</html>
"""


def _extract_intro_text(step_summaries: list[dict]) -> str:
    for item in reversed(step_summaries):
        content = _extract_summary_content(item.get("content", ""))
        if not content:
            continue
        match = re.search(r"执行结果[：:]\s*([^\n]+)", content)
        if match:
            return match.group(1).strip()
        return content.splitlines()[0].strip()
    return "已基于分步骤成果重组最终报告页面。"


def _normalize_section_title(title: str) -> str:
    normalized = re.sub(r"[：:].*$", "", title or "").strip()
    return normalized or "报告章节"


def _render_summary_html(text: str) -> str:
    normalized = format_summary_for_display(text)
    if not normalized:
        return '<div class="rich-text"></div>'

    blocks = [block.strip() for block in re.split(r"\n\s*\n", normalized) if block.strip()]
    rendered_blocks: list[str] = []
    for block in blocks:
        lines = [line.strip() for line in block.splitlines() if line.strip()]
        list_items = []
        paragraph_lines = []
        for line in lines:
            cleaned_line = re.sub(r"^[\-•●]\s*", "", line)
            cleaned_line = re.sub(r"^\d+[\.、]\s*", "", cleaned_line)
            if line.startswith(("-", "•", "●")) or re.match(r"^\d+[\.、]", line):
                list_items.append(f"<li>{_render_inline_markdown_to_html(cleaned_line)}</li>")
            else:
                paragraph_lines.append(_render_inline_markdown_to_html(line))

        if paragraph_lines:
            paragraph_html = "<br/>".join(paragraph_lines)
            css_class = "callout" if any(token in block for token in ["执行结果", "核心洞察", "策略建议"]) else ""
            if css_class:
                rendered_blocks.append(f'<div class="{css_class}">{paragraph_html}</div>')
            else:
                rendered_blocks.append(f"<p>{paragraph_html}</p>")
        if list_items:
            rendered_blocks.append(f"<ul>{''.join(list_items)}</ul>")

    return f'<div class="rich-text">{"".join(rendered_blocks)}</div>'


def _build_structured_delivery_text(context: dict) -> str:
    title = str(context.get("title") or "交付结果").replace("_", " ")
    intro_text = str(context.get("intro_text") or "").strip()
    sections = [f"# {title}"]
    if intro_text:
        sections.append(f"**执行摘要：**\n{format_summary_for_display(intro_text)}")

    for item in context.get("step_summaries") or []:
        content = format_summary_for_display(item.get("content", ""))
        if not content:
            continue
        sections.append(f"## {_normalize_section_title(item.get('title') or '')}\n{content}")

    return "\n\n".join(sections).strip() + "\n"


def _build_markdown_document(title: str, body_text: str, preserve_body: bool = False) -> str:
    if preserve_body:
        return (body_text or "").strip() + "\n"

    formatted_body = format_summary_for_display(body_text)
    if re.search(r"^#{1,6}\s", formatted_body, flags=re.MULTILINE):
        return formatted_body
    return f"# {title}\n\n{formatted_body.strip()}\n"


def _build_pdf_document(file_path: Path, title: str, body_text: str) -> None:
    if SimpleDocTemplate is None or registerFont is None or UnicodeCIDFont is None:
        raise HTTPException(status_code=500, detail="PDF generation dependency is not installed")

    registerFont(UnicodeCIDFont("STSong-Light"))
    buffer = BytesIO()
    document = SimpleDocTemplate(buffer, pagesize=A4, leftMargin=48, rightMargin=48, topMargin=54, bottomMargin=54)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("ArtifactTitle", parent=styles["Heading1"], fontName="STSong-Light", fontSize=20, leading=28, spaceAfter=16)
    body_style = ParagraphStyle("ArtifactBody", parent=styles["BodyText"], fontName="STSong-Light", fontSize=11, leading=20, spaceAfter=10)

    story = [Paragraph(escape(title), title_style), Spacer(1, 8)]
    for paragraph in [segment.strip() for segment in re.split(r"\n\s*\n", body_text) if segment.strip()]:
        story.append(Paragraph(escape(paragraph).replace("\n", "<br/>"), body_style))
        story.append(Spacer(1, 6))

    document.build(story)
    file_path.write_bytes(buffer.getvalue())


def _build_pptx_document(file_path: Path, title: str, body_text: str) -> None:
    if Presentation is None or Pt is None:
        raise HTTPException(status_code=500, detail="PPTX generation dependency is not installed")

    presentation = Presentation()
    title_slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    title_slide.shapes.title.text = title
    title_slide.placeholders[1].text = "根据本轮对话自动生成的演示文稿"

    readable_body = _prepare_pptx_content(body_text)
    sections = [segment.strip() for segment in re.split(r"\n\s*\n|(?=第[一二三四五六七八九十\d]+[章节部分])", readable_body) if segment.strip()]
    for index, section in enumerate(sections[:8], start=1):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = f"第 {index} 部分"
        text_frame = slide.placeholders[1].text_frame
        text_frame.clear()

        lines = [line.strip() for line in section.splitlines() if line.strip()] or [section]
        for line_index, line in enumerate(lines[:6]):
            paragraph = text_frame.paragraphs[0] if line_index == 0 else text_frame.add_paragraph()
            paragraph.text = line
            paragraph.level = 0
            for run in paragraph.runs:
                run.font.size = Pt(20 if line_index == 0 else 16)

    presentation.save(str(file_path))