import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch


BACKEND_DIR = Path(__file__).resolve().parent
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))


from app.services import artifact_service


class ArtifactServiceHtmlTests(unittest.TestCase):
    def test_html_artifact_excludes_surrounding_chat_text(self):
        response_text = """由于终端指令受限，我直接将 HTML-PPT 的内容呈现在下方。

```html
<!DOCTYPE html>
<html lang="zh-CN">
<head><title>机器人调研报告</title></head>
<body><main>报告内容</main></body>
</html>
```

请将以上内容保存为 report.html。
"""

        with tempfile.TemporaryDirectory() as temp_dir:
            project_root = Path(temp_dir)
            workspace_root = project_root / "workspace"
            artifact_root = workspace_root / "generated"
            artifact_root.mkdir(parents=True, exist_ok=True)

            with patch.object(artifact_service, "PROJECT_ROOT", project_root), \
                 patch.object(artifact_service, "WORKSPACE_ROOT", workspace_root), \
                 patch.object(artifact_service, "ARTIFACT_ROOT", artifact_root):
                artifacts = artifact_service.create_artifacts_for_response(
                    user_request="请生成 HTML-PPT 给我",
                    response_text=response_text,
                )

            html_artifacts = [artifact for artifact in artifacts if artifact["format"] == "html"]
            self.assertEqual(len(html_artifacts), 1)
            file_path = project_root / html_artifacts[0]["relative_path"]
            content = file_path.read_text(encoding="utf-8")
            self.assertTrue(content.startswith("<!DOCTYPE html>"))
            self.assertIn("<main>报告内容</main>", content)
            self.assertNotIn("终端指令受限", content)
            self.assertNotIn("```", content)
            self.assertNotIn("请将以上内容保存", content)

    def test_pptx_artifact_extracts_readable_text_from_embedded_html(self):
        response_text = """由于当前环境限制，以下是 HTML PPT 源码。\n+
```html
<!DOCTYPE html>
<html><head><title>核心价值总结</title><style>.hidden{display:none}</style></head>
<body><section><h1>核心价值总结</h1><p>帮助团队理解业务对象。</p><ul><li>统一数据语义</li><li>提升决策效率</li></ul></section>
<script>private_code()</script></body></html>
```
"""

        with tempfile.TemporaryDirectory() as temp_dir:
            project_root = Path(temp_dir)
            workspace_root = project_root / "workspace"
            artifact_root = workspace_root / "generated"
            artifact_root.mkdir(parents=True, exist_ok=True)

            with patch.object(artifact_service, "PROJECT_ROOT", project_root), \
                 patch.object(artifact_service, "WORKSPACE_ROOT", workspace_root), \
                 patch.object(artifact_service, "ARTIFACT_ROOT", artifact_root):
                artifacts = artifact_service.create_artifacts_for_response(
                    user_request="请生成 PPTX",
                    response_text=response_text,
                )

            pptx_artifact = next(artifact for artifact in artifacts if artifact["format"] == "pptx")
            pptx_path = project_root / pptx_artifact["relative_path"]
            from pptx import Presentation

            slide_text = "\n".join(
                shape.text
                for slide in Presentation(str(pptx_path)).slides
                for shape in slide.shapes
                if hasattr(shape, "text")
            )
            self.assertIn("帮助团队理解业务对象", slide_text)
            self.assertIn("统一数据语义", slide_text)
            self.assertNotIn("private_code", slide_text)
            self.assertNotIn("当前环境限制", slide_text)


if __name__ == "__main__":
    unittest.main()