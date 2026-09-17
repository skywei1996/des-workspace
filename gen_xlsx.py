# coding: utf-8
import openpyxl, os
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill

output_dir = "D:/cursor/DES workspace - Copy/12.12 - 公司最终版本/workspace/generated/codex_277af18f4a5f4d1e89a7b6201516a405"
os.makedirs(output_dir, exist_ok=True)
output_path = os.path.join(output_dir, "调补货执行结果.xlsx")

wb = openpyxl.Workbook()

# ===== Sheet 1 =====
ws = wb.active
ws.title = "调补货执行结果"

headers = ["序号","门店ID","门店名称","城市","SKU","调入数量","可用库存","近7天销量","近14天销量","可销售周数","需审批","决策原因"]

hfont = Font(name="Arial", bold=True, size=11, color="FFFFFF")
hfill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
halign = Alignment(horizontal="center", vertical="center", wrap_text=True)
border = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin")
)
dfont = Font(name="Arial", size=10)
dalign = Alignment(horizontal="center", vertical="center", wrap_text=True)

for i,h in enumerate(headers, 1):
    c = ws.cell(row=1, column=i, value=h)
    c.font = hfont; c.fill = hfill; c.alignment = halign; c.border = border

rows = [
    [1,"ZJ-HUZ-001","湖州湖州银泰城专柜","湖州","W26AU0001-01-M",5,4,8,14,0.57,"否","近7天有销量(8件)，可销售周数<1，需要补货"],
    [2,"ZJ-HUZ-001","湖州湖州银泰城专柜","湖州","W26AU0001-01-S",4,4,2,12,0.67,"否","近7天有销量(2件)，可销售周数<1，需要补货"],
    [3,"ZJ-HUZ-001","湖州湖州银泰城专柜","湖州","W26AU0001-02-M",6,9,8,25,0.72,"否","近7天有销量(8件)，可销售周数<1，需要补货"],
]
for ri,row in enumerate(rows, 2):
    for ci,val in enumerate(row, 1):
        c = ws.cell(row=ri, column=ci, value=val)
        c.font = dfont; c.alignment = dalign; c.border = border

widths = [6,14,26,10,20,10,10,12,12,12,8,44]
for i,w in enumerate(widths, 1):
    ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

# ===== Sheet 2: 校验 =====
ws2 = wb.create_sheet("可销售周数校验")
h2 = ["序号","SKU","可用库存","近14天销量","近7天销量","平均周销量(14天)","可销售周数(系统)","可销售周数(校验)","校验结果"]
for i,h in enumerate(h2, 1):
    c = ws2.cell(row=1, column=i, value=h)
    c.font = hfont; c.fill = hfill; c.alignment = halign; c.border = border

vdata = [
    [1,"W26AU0001-01-M",4,14,8,7.0,0.57,round(4/7,2),"一致"],
    [2,"W26AU0001-01-S",4,12,2,6.0,0.67,round(4/6,2),"一致"],
    [3,"W26AU0001-02-M",9,25,8,12.5,0.72,round(9/12.5,2),"一致"],
]
for ri,row in enumerate(vdata, 2):
    for ci,val in enumerate(row, 1):
        c = ws2.cell(row=ri, column=ci, value=val)
        c.font = dfont; c.alignment = dalign; c.border = border

w2 = [6,20,10,12,12,16,16,16,10]
for i,w in enumerate(w2, 1):
    ws2.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

# ===== Sheet 3: 说明 =====
ws3 = wb.create_sheet("执行说明")
h3 = ["项目","内容"]
for i,h in enumerate(h3, 1):
    c = ws3.cell(row=1, column=i, value=h)
    c.font = hfont; c.fill = hfill; c.alignment = halign; c.border = border

notes = [
    ["执行时间","2026-09-16"],
    ["项目ID","0b1a25a9-f6b7-4f38-8007-38d4182e3e4a"],
    ["门店","ZJ-HUZ-001（湖州湖州银泰城专柜）"],
    ["调补货方向","调入（inbound）"],
    ["目标可销售周数","1.2"],
    ["API状态","数据映射未配置（Mapped dataset not found for 销售明细），使用用户提供的数据"],
    ["SKU数量","3个"],
    ["需审批条目","0条"],
    ["总体状态","数据已就绪，待数据映射配置完成后可通过API完整执行"],
    ["可销售周数校验","所有3条记录的WOS计算均正确，公式为：可用库存 / (近14天销量/2)"],
]
for ri,(k,v) in enumerate(notes, 2):
    c1 = ws3.cell(row=ri, column=1, value=k); c1.font = dfont; c1.alignment = dalign; c1.border = border
    c2 = ws3.cell(row=ri, column=2, value=v); c2.font = dfont; c2.alignment = Alignment(vertical="center", wrap_text=True); c2.border = border

ws3.column_dimensions["A"].width = 20
ws3.column_dimensions["B"].width = 80

wb.save(output_path)
print("OK saved")
wb2 = openpyxl.load_workbook(output_path)
print("Sheets:", wb2.sheetnames)
print("Total rows:", sum(s.max_row for s in wb2.worksheets))
